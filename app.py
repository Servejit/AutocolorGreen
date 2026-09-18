import streamlit as st
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from copy import copy
import re
import io


# ============================================================
# PAGE
# ============================================================

st.set_page_config(
    page_title="6thsense Vardaan Auto Color",
    layout="wide"
)

st.title("6thsense Vardaan Auto Color")


# ============================================================
# HELPERS
# ============================================================

def clean_text(value):
    if value is None:
        return ""
    text = str(value)
    text = text.replace("\u00a0", " ")
    text = text.replace("\u200b", "")
    text = text.replace("\u200c", "")
    text = text.replace("\u200d", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip().lower()


def get_number(value):
    if value is None:
        return None

    if isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    text = text.replace(",", "").replace("%", "")

    try:
        return float(text)
    except:
        return None


def get_parentheses_number(value):
    if value is None:
        return None

    text = str(value)

    match = re.search(
        r"\(\s*([-+]?\d+(?:\.\d+)?)\s*\)",
        text
    )

    if match:
        try:
            return float(match.group(1))
        except:
            return None

    return None


# ============================================================
# FILE UPLOAD
# ============================================================

uploaded_file = st.file_uploader(
    "Upload Excel file",
    type=["xlsx", "xlsm"]
)


if uploaded_file is None:
    st.info("Upload your master Excel file to continue.")
    st.stop()


# ============================================================
# LOAD ORIGINAL WORKBOOK
# ============================================================

file_bytes = uploaded_file.getvalue()

value_wb = openpyxl.load_workbook(
    io.BytesIO(file_bytes),
    data_only=True
)

format_wb = openpyxl.load_workbook(
    io.BytesIO(file_bytes),
    data_only=False
)


# ============================================================
# FIND SUMMARY SHEET
# ============================================================

if "summary" not in value_wb.sheetnames:
    st.error("Sheet named 'summary' was not found.")
    st.stop()

value_ws = value_wb["summary"]
format_ws = format_wb["summary"]


# ============================================================
# CREATE OUTPUT WORKBOOK
# ============================================================

out_wb = Workbook()

out_ws = out_wb.active
out_ws.title = "summary"


# ============================================================
# COPY VALUES + FORMATTING
# ============================================================

for row in value_ws.iter_rows():

    for value_cell in row:

        r = value_cell.row
        c = value_cell.column

        out_cell = out_ws.cell(
            row=r,
            column=c,
            value=value_cell.value
        )

        source_cell = format_ws.cell(
            row=r,
            column=c
        )

        # Font
        if source_cell.has_style:
            out_cell.font = copy(source_cell.font)
            out_cell.fill = copy(source_cell.fill)
            out_cell.border = copy(source_cell.border)
            out_cell.alignment = copy(source_cell.alignment)
            out_cell.protection = copy(source_cell.protection)

        # Number format
        out_cell.number_format = source_cell.number_format


# ============================================================
# COPY COLUMN WIDTHS
# ============================================================

for key, dimension in format_ws.column_dimensions.items():

    out_ws.column_dimensions[key].width = dimension.width
    out_ws.column_dimensions[key].hidden = dimension.hidden


# ============================================================
# COPY ROW HEIGHTS
# ============================================================

for key, dimension in format_ws.row_dimensions.items():

    out_ws.row_dimensions[key].height = dimension.height
    out_ws.row_dimensions[key].hidden = dimension.hidden


# ============================================================
# COPY MERGED CELLS
# ============================================================

for merged_range in format_ws.merged_cells.ranges:
    out_ws.merge_cells(str(merged_range))


# ============================================================
# FIND HEADER ROW
# ============================================================

header_row = None

for row in range(1, out_ws.max_row + 1):

    value = out_ws.cell(row, 1).value

    if clean_text(value) == "symbol":
        header_row = row
        break


if header_row is None:
    st.error("Could not find header row containing 'Symbol' in Column A.")
    st.stop()


# ============================================================
# BLUE FILL
# ============================================================

blue_fill = PatternFill(
    fill_type="solid",
    fgColor="ADD8E6"
)

blue_cells = set()


def make_blue(row, col):

    out_ws.cell(row, col).fill = copy(blue_fill)

    blue_cells.add((row, col))

    # Corresponding Column A
    out_ws.cell(row, 1).fill = copy(blue_fill)


# ============================================================
# FIND HEADINGS
# ============================================================

headings = {}

for col in range(1, out_ws.max_column + 1):

    heading = clean_text(
        out_ws.cell(header_row, col).value
    )

    if heading:
        headings[heading] = col


# ============================================================
# RULE 1
# Sum I < -4
# ============================================================

sum_i_col = None

for col in range(1, out_ws.max_column + 1):

    heading = clean_text(
        out_ws.cell(header_row, col).value
    )

    if heading == "sum i":
        sum_i_col = col
        break


if sum_i_col:

    for row in range(header_row + 1, out_ws.max_row + 1):

        number = get_number(
            out_ws.cell(row, sum_i_col).value
        )

        if number is not None and number < -4:
            make_blue(row, sum_i_col)


# ============================================================
# RULE 2
# 16> C-B / Avg.4
# Parentheses value < 0.50
# ============================================================

col_c = None

for col in range(1, out_ws.max_column + 1):

    heading = clean_text(
        out_ws.cell(header_row, col).value
    )

    if heading == "16> c-b / avg.4":
        col_c = col
        break


if col_c:

    for row in range(header_row + 1, out_ws.max_row + 1):

        value = get_parentheses_number(
            out_ws.cell(row, col_c).value
        )

        if value is not None and value < 0.50:
            make_blue(row, col_c)


# ============================================================
# RULE 3
# 16< D-B / Avg.4
#
# Parentheses value < -1
# AND
# Parentheses value < number immediately after 16<
# ============================================================

col_d = None

for col in range(1, out_ws.max_column + 1):

    heading = clean_text(
        out_ws.cell(header_row, col).value
    )

    if heading == "16< d-b / avg.4":
        col_d = col
        break


if col_d:

    for row in range(header_row + 1, out_ws.max_row + 1):

        cell_value = out_ws.cell(row, col_d).value

        if cell_value is None:
            continue

        text = str(cell_value)

        first_match = re.search(
            r"16<\s*([-+]?\d+(?:\.\d+)?)",
            text,
            re.IGNORECASE
        )

        parent_value = get_parentheses_number(
            cell_value
        )

        if first_match and parent_value is not None:

            try:
                changing_value = float(
                    first_match.group(1)
                )
            except:
                continue

            if (
                parent_value < -1
                and parent_value < changing_value
            ):
                make_blue(row, col_d)


# ============================================================
# RULE 4
# Sum O2H.10
#
# Values strictly below average
# ============================================================

sum_o2h_col = None

for col in range(1, out_ws.max_column + 1):

    heading = clean_text(
        out_ws.cell(header_row, col).value
    )

    if heading == "sum o2h.10":
        sum_o2h_col = col
        break


if sum_o2h_col:

    values = []

    for row in range(header_row + 1, out_ws.max_row + 1):

        number = get_number(
            out_ws.cell(row, sum_o2h_col).value
        )

        if number is not None:
            values.append(number)

    if values:

        average_value = sum(values) / len(values)

        for row in range(header_row + 1, out_ws.max_row + 1):

            number = get_number(
                out_ws.cell(row, sum_o2h_col).value
            )

            if number is not None and number < average_value:
                make_blue(row, sum_o2h_col)


# ============================================================
# RULE 5
# Sum O2L.10
#
# Values strictly below average
# ============================================================

sum_o2l_col = None

for col in range(1, out_ws.max_column + 1):

    heading = clean_text(
        out_ws.cell(header_row, col).value
    )

    if heading == "sum o2l.10":
        sum_o2l_col = col
        break


if sum_o2l_col:

    values = []

    for row in range(header_row + 1, out_ws.max_row + 1):

        number = get_number(
            out_ws.cell(row, sum_o2l_col).value
        )

        if number is not None:
            values.append(number)

    if values:

        average_value = sum(values) / len(values)

        for row in range(header_row + 1, out_ws.max_row + 1):

            number = get_number(
                out_ws.cell(row, sum_o2l_col).value
            )

            if number is not None and number < average_value:
                make_blue(row, sum_o2l_col)


# ============================================================
# RULE 6
# FIRST TWO DATED O2L COLUMNS
# Excluding Sum O2L.10
# Value < -1
# ============================================================

dated_o2l_columns = []

for col in range(1, out_ws.max_column + 1):

    heading = clean_text(
        out_ws.cell(header_row, col).value
    )

    if (
        "o2l" in heading
        and heading != "sum o2l.10"
    ):
        dated_o2l_columns.append(col)


for col in dated_o2l_columns[:2]:

    for row in range(header_row + 1, out_ws.max_row + 1):

        number = get_number(
            out_ws.cell(row, col).value
        )

        if number is not None and number < -1:
            make_blue(row, col)


# ============================================================
# RULE 7
# 10 "-ve"
#
# Example:
# 2+1,2,3,4,6
#
# Check ONLY first number before +
# ============================================================

negative_col = None

for col in range(1, out_ws.max_column + 1):

    heading = clean_text(
        out_ws.cell(header_row, col).value
    )

    if heading == '10 "-ve"':
        negative_col = col
        break


if negative_col:

    for row in range(header_row + 1, out_ws.max_row + 1):

        value = out_ws.cell(row, negative_col).value

        if value is None:
            continue

        text = str(value).strip()

        match = re.match(
            r"\s*(\d+)\s*\+",
            text
        )

        if match:

            try:
                first_number = int(match.group(1))
            except:
                continue

            if first_number > 1:
                make_blue(row, negative_col)


# ============================================================
# RULE 8
# %Chg.1 < 0
# ============================================================

pct1_col = None

for col in range(1, out_ws.max_column + 1):

    heading = clean_text(
        out_ws.cell(header_row, col).value
    )

    if heading == "%chg.1":
        pct1_col = col
        break


if pct1_col:

    for row in range(header_row + 1, out_ws.max_row + 1):

        number = get_number(
            out_ws.cell(row, pct1_col).value
        )

        if number is not None and number < 0:
            make_blue(row, pct1_col)


# ============================================================
# RULE 9
# %Chg.2 < 0
# ============================================================

pct2_col = None

for col in range(1, out_ws.max_column + 1):

    heading = clean_text(
        out_ws.cell(header_row, col).value
    )

    if heading == "%chg.2":
        pct2_col = col
        break


if pct2_col:

    for row in range(header_row + 1, out_ws.max_row + 1):

        number = get_number(
            out_ws.cell(row, pct2_col).value
        )

        if number is not None and number < 0:
            make_blue(row, pct2_col)


# ============================================================
# COUNT BLUE CELLS PER ROW
#
# Column A is NOT counted.
# ============================================================

row_blue_counts = {}

for row in range(header_row + 1, out_ws.max_row + 1):

    count = sum(
        1
        for r, c in blue_cells
        if r == row and c != 1
    )

    row_blue_counts[row] = count


# ============================================================
# THREE GREEN LEVELS
#
# Highest unique blue count = DARK GREEN
# Second highest = GREEN
# Third highest = LIGHT GREEN
#
# Only rows with >= 8 blue cells qualify.
# Only Column A receives green.
# ============================================================

dark_green_fill = PatternFill(
    fill_type="solid",
    fgColor="548235"
)

green_fill = PatternFill(
    fill_type="solid",
    fgColor="70AD47"
)

light_green_fill = PatternFill(
    fill_type="solid",
    fgColor="90EE90"
)


eligible_counts = sorted(
    {
        count
        for count in row_blue_counts.values()
        if count >= 8
    },
    reverse=True
)


green_counts = eligible_counts[:3]


for row, count in row_blue_counts.items():

    if count < 8:
        continue

    if len(green_counts) >= 1 and count == green_counts[0]:

        out_ws.cell(row, 1).fill = copy(
            dark_green_fill
        )

    elif len(green_counts) >= 2 and count == green_counts[1]:

        out_ws.cell(row, 1).fill = copy(
            green_fill
        )

    elif len(green_counts) >= 3 and count == green_counts[2]:

        out_ws.cell(row, 1).fill = copy(
            light_green_fill
        )


# ============================================================
# AUTOFILTER
# ============================================================

out_ws.auto_filter.ref = (
    f"A{header_row}:"
    f"{openpyxl.utils.get_column_letter(out_ws.max_column)}"
    f"{out_ws.max_row}"
)


# ============================================================
# FREEZE HEADER
# ============================================================

out_ws.freeze_panes = f"A{header_row + 1}"


# ============================================================
# SAVE TO MEMORY
# ============================================================

output_buffer = io.BytesIO()

out_wb.save(output_buffer)

output_buffer.seek(0)


# ============================================================
# DOWNLOAD
# ============================================================

st.success("Excel file processed successfully.")

st.download_button(
    label="Download 6thsenseVardaanAutocolor.xlsx",
    data=output_buffer.getvalue(),
    file_name="6thsenseVardaanAutocolor.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
